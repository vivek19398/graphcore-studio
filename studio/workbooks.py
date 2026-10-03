"""Bounded, read-only XLSX import into the shared typed artifact store."""
import base64
import binascii
import csv
import datetime
import hashlib
import io
import zipfile
from openpyxl import load_workbook

class WorkbookImport:
    MAX_BYTES = 2 * 1024 * 1024

    def __init__(self, tables):
        self.tables = tables

    def _decode(self, name, content):
        if not isinstance(name,str) or not name.lower().endswith('.xlsx') or '/' in name or '\\' in name or len(name)>200:
            raise ValueError('Choose a .xlsx filename without directory paths; .xls and macros are unsupported')
        if not isinstance(content,str) or len(content)>4*((self.MAX_BYTES+2)//3):
            raise ValueError('Workbook must be up to 2 MiB')
        try: raw=base64.b64decode(content,validate=True)
        except (ValueError,binascii.Error): raise ValueError('Invalid workbook encoding')
        if len(raw)>self.MAX_BYTES: raise ValueError('Workbook exceeds 2 MiB')
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries=archive.infolist()
                if len(entries)>1000 or sum(e.file_size for e in entries)>16*1024*1024 or any(e.file_size>8*1024*1024 for e in entries):
                    raise ValueError('Workbook expands beyond the supported local import limits')
                if any('vbaproject' in e.filename.lower() for e in entries):
                    raise ValueError('Macro-enabled workbooks are unsupported')
                if any(e.flag_bits & 1 for e in entries): raise ValueError('Encrypted workbooks are unsupported')
        except zipfile.BadZipFile: raise ValueError('Invalid XLSX archive')
        return raw

    def _load(self, raw, data_only):
        try: return load_workbook(io.BytesIO(raw),read_only=True,data_only=data_only,keep_links=False)
        except Exception as error: raise ValueError('Cannot read XLSX workbook: '+str(error)) from error

    def sheets(self,name,content):
        raw=self._decode(name,content)
        workbook=self._load(raw,False)
        try:
            sheets=[{'name':s.title,'rows':s.max_row,'columns':s.max_column} for s in workbook.worksheets if s.sheet_state=='visible']
            if not 1<=len(sheets)<=50: raise ValueError('Workbook needs 1–50 visible worksheets')
            return {'name':name,'sheets':sheets}
        finally: workbook.close()

    def _extract(self,name,content,sheet,header_row):
        raw=self._decode(name,content)
        if isinstance(header_row,bool) or not isinstance(header_row,int) or not 1<=header_row<=100:
            raise ValueError('Header row must be between 1 and 100')
        formulas=self._load(raw,False)
        values=None
        try:
            if not isinstance(sheet,str) or sheet not in formulas.sheetnames or formulas[sheet].sheet_state!='visible':
                raise ValueError('Choose a visible worksheet')
            fs=formulas[sheet]
            if (fs.max_column or 0)>self.tables.MAX_COLUMNS or (fs.max_row or 0)>header_row+self.tables.MAX_ROWS:
                raise ValueError('Worksheet exceeds 100 columns or 20000 data rows; remove unused formatted ranges if needed')
            values=self._load(raw,True)
            vs=values[sheet]
            # Do not trust understated worksheet dimension metadata.
            fs.reset_dimensions(); vs.reset_dimensions()
            formula_rows=fs.iter_rows(min_row=header_row,max_col=self.tables.MAX_COLUMNS+1)
            value_rows=vs.iter_rows(min_row=header_row,max_col=self.tables.MAX_COLUMNS+1)
            records=[]; occupied=[]; formula_count=missing=0
            for index,(frow,vrow) in enumerate(zip(formula_rows,value_rows)):
                if index>self.tables.MAX_ROWS: raise ValueError("Worksheet exceeds 20000 data rows")
                present=[i for i,c in enumerate(frow) if c.value is not None]
                if index==0:
                    if not present: raise ValueError('Selected header row is empty')
                    width=max(present)+1
                    if width>self.tables.MAX_COLUMNS: raise ValueError('Worksheet exceeds 100 columns')
                    if any(c.data_type=='f' for c in frow[:width]): raise ValueError('Column headers must not contain formulas')
                elif any(i>=width for i in present): raise ValueError('Worksheet has data outside the header columns')
                converted=[]
                for fcell,vcell in zip(frow[:width],vrow[:width]):
                    if fcell.data_type=='f':
                        formula_count+=1
                        if vcell.value is None: missing+=1
                    if vcell.data_type=='e': raise ValueError('Worksheet contains an Excel error at '+vcell.coordinate)
                    value=vcell.value
                    if isinstance(value,bool): value='true' if value else 'false'
                    elif isinstance(value,(datetime.datetime,datetime.date,datetime.time)): value=value.isoformat()
                    converted.append('' if value is None else str(value))
                records.append(converted)
                occupied.append(bool(present))
            if not records: raise ValueError('Worksheet has no header row')
            # Ignore trailing blank rows, preserving internal blank rows as null records.
            while len(records)>1 and not occupied[-1]:
                records.pop(); occupied.pop()
            if len(records)-1>self.tables.MAX_ROWS: raise ValueError('Worksheet exceeds 20000 data rows')
            wire=io.StringIO();csv.writer(wire).writerows(records)
            warnings=['Excel numeric values use the precision stored in the workbook. Dates become ISO text; display formatting is not imported.']
            if formula_count: warnings.append('%d formula cells use cached values; freshness cannot be verified. No formulas are calculated.'%formula_count)
            if missing: warnings.append('%d formula cells have no cached result. Recalculate and save the workbook in Excel before importing.'%missing)
            return raw,wire.getvalue(),warnings,formula_count,missing
        finally:
            formulas.close()
            if values is not None: values.close()

    def inspect(self,name,content,sheet,header_row=1):
        raw,text,warnings,count,missing=self._extract(name,content,sheet,header_row)
        preview=self.tables.inspect('worksheet.csv',text)
        return dict(preview,name=name,warnings=warnings,formula_count=count,missing_formula_results=missing)

    def add(self,name,content,sheet,header_row=1,types=None,acknowledge_formulas=False):
        raw,text,warnings,count,missing=self._extract(name,content,sheet,header_row)
        if missing: raise ValueError('Formula results are missing. Recalculate and save in Excel, then choose the workbook again')
        if count and acknowledge_formulas is not True: raise ValueError('Acknowledge that formula caches may be stale before importing')
        table=self.tables._parse('worksheet.csv',text,types)
        table.update(name=name,source_sha256=hashlib.sha256(raw).hexdigest(),
                     provenance={'operation':'xlsx_import','sheet':sheet,'header_row':header_row,'formula_count':count,'warnings':warnings})
        return self.tables._persist(table)
