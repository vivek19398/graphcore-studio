import base64
import io
import tempfile
import unittest
import zipfile
from openpyxl import Workbook
from studio.tables import LocalTables
from studio.workbooks import WorkbookImport

def fixture(cached=False):
    wb=Workbook(); ws=wb.active;ws.title='Sales'
    ws.append(['Report']);ws.append(['team','amount','active'])
    ws.append(['A',0.1,True]);ws.append(['A',0.2,False])
    formulas=wb.create_sheet('Formulas');formulas.append(['amount']);formulas.append(['=1+2'])
    hidden=wb.create_sheet('Hidden');hidden.sheet_state='hidden'
    stream=io.BytesIO();wb.save(stream);raw=stream.getvalue()
    if cached:
        source=zipfile.ZipFile(io.BytesIO(raw));output=io.BytesIO()
        with zipfile.ZipFile(output,'w') as dest:
            for entry in source.infolist():
                content=source.read(entry.filename)
                if entry.filename=='xl/worksheets/sheet2.xml': content=content.replace(b'<v></v>',b'<v>3</v>').replace(b'<v />',b'<v>3</v>')
                dest.writestr(entry,content)
        source.close();raw=output.getvalue()
    return base64.b64encode(raw).decode()

class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.tables=LocalTables(self.temp.name);self.importer=WorkbookImport(self.tables)

    def test_sheet_header_typed_import_and_source_hash(self):
        content=fixture()
        self.assertEqual([s['name'] for s in self.importer.sheets('sales.xlsx',content)['sheets']],['Sales','Formulas'])
        preview=self.importer.inspect('sales.xlsx',content,'Sales',2)
        self.assertEqual(preview['row_count'],2);self.assertEqual(self.tables.list(),[])
        ref=self.importer.add('sales.xlsx',content,'Sales',2,{'amount':'decimal','active':'boolean'})
        self.assertEqual(self.tables.summarize(ref,'amount')['sum'],'0.3')
        self.assertEqual(ref['name'],'sales.xlsx');self.assertEqual(ref['provenance']['header_row'],2)
        self.assertIs(self.tables.read(ref['id'])['rows'][0]['active'],True)

    def test_formula_warnings_acknowledgment_and_missing_cache(self):
        uncached=fixture()
        preview=self.importer.inspect('sales.xlsx',uncached,'Formulas')
        self.assertEqual(preview['missing_formula_results'],1)
        self.assertEqual(preview['row_count'],1)
        with self.assertRaisesRegex(ValueError,'missing'): self.importer.add('sales.xlsx',uncached,'Formulas',types={'amount':'decimal'},acknowledge_formulas=True)
        cached=fixture(True)
        with self.assertRaisesRegex(ValueError,'Acknowledge'): self.importer.add('sales.xlsx',cached,'Formulas')
        ref=self.importer.add('sales.xlsx',cached,'Formulas',types={'amount':'decimal'},acknowledge_formulas=True)
        self.assertEqual(self.tables.summarize(ref,'amount')['sum'],'3')
        self.assertEqual(ref['provenance']['formula_count'],1)

    def test_invalid_archive_sheet_headers_and_encoding(self):
        content=fixture()
        for name,wire,sheet,header in [('bad.xls',content,'Sales',2),('bad.xlsx','!!','Sales',2),('sales.xlsx',content,'Hidden',1),('sales.xlsx',content,'Sales',0),('sales.xlsx',content,'Sales',1)]:
            with self.subTest(name=name,sheet=sheet,header=header),self.assertRaises(ValueError):
                self.importer.inspect(name,wire,sheet,header)
        self.assertEqual(self.tables.list(),[])

    def test_archive_expansion_and_macro_rejection(self):
        for filename,payload in [('xl/vbaProject.bin',b'macro'),('xl/worksheets/sheet1.xml',b'x'*(9*1024*1024))]:
            stream=io.BytesIO()
            with zipfile.ZipFile(stream,'w',compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(filename,payload)
            with self.subTest(filename=filename),self.assertRaises(ValueError):
                self.importer.sheets('bad.xlsx',base64.b64encode(stream.getvalue()).decode())

    def test_dates_and_header_validation(self):
        import datetime
        wb=Workbook();ws=wb.active;ws.append(['date','amount']);ws.append([datetime.date(2026,10,3),1])
        stream=io.BytesIO();wb.save(stream)
        content=base64.b64encode(stream.getvalue()).decode()
        preview=self.importer.inspect('dates.xlsx',content,'Sheet')
        self.assertTrue(preview['rows'][0]['date'].startswith('2026-10-03'))
        ws.cell(1,2,'date');stream=io.BytesIO();wb.save(stream)
        with self.assertRaisesRegex(ValueError,'unique'):
            self.importer.inspect('duplicates.xlsx',base64.b64encode(stream.getvalue()).decode(),'Sheet')
